#!/usr/bin/env python3
"""City offtake, store list and tentative beat options for the corporate visit (Excel) + a small city summary CSV.

Reads the maintained store master (PowerBI/SeedData/Masters/Store_City_Master.csv, built by build_store_city_master.py)
and the monthly store x article offtake files. The Excel goes wherever --out says and is NOT committed (it lists stores).
The city summary CSV (no store rows) is written to data/nielsen/Visit_City_Summary_<label>.csv for the dashboards.

    python scripts/build_visit_city_plan.py --out City_Offtake_Beat_Plan.xlsx --label Aug26 [--months Apr May Jun Jul Aug]

Beats: 5 for each of the 8 biggest listed cities by Apr-Aug NSV, 3 for the rest. Three options per store (A chain-wise,
B sales tiers, C mixed). There is no address or pincode in the masters, so a beat is a group of stores, not a route;
GT and EBO teams add the area and pick the final beat.
Last-year store sales and stock / SOS are not in the repo: those columns stay blank and are named in the README.
"""
import argparse
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

import visit_cities as vc

ROOT = vc.ROOT
MASTER = ROOT / "PowerBI" / "SeedData" / "Masters" / "Store_City_Master.csv"
LY_JSON = ROOT / "data" / "raw_drops" / "_agg" / "offtake_fy26.json"
LYM = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25"]
HF = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
GFILL = PatternFill("solid", fgColor="BFBFBF")
YFILL = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="E2EFDA")
RED = PatternFill("solid", fgColor="FCE4D6")
AMBER = PatternFill("solid", fgColor="FFE699")


def fv(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    return round(float(v), 4) if isinstance(v, (int, float)) else v


def header(ws, row, cols, grey=()):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c)
        x.font = HF if c not in grey else Font(bold=True)
        x.fill = HFILL if c not in grey else GFILL
        x.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def widths(ws, w):
    for i, v in enumerate(w, 1):
        ws.column_dimensions[L(i)].width = v


def balanced(df, n, by):
    load, out = [0.0] * n, {}
    for k, v in df.groupby(by)["Total"].sum().sort_values(ascending=False).items():
        i = load.index(min(load))
        out[k] = i + 1
        load[i] += v
    return out


def opt_chain(df, n):
    return df["Chain"].map(balanced(df, n, "Chain"))


def opt_tier(df, n):
    x = df.sort_values("Total", ascending=False)
    cs = x["Total"].fillna(0).cumsum() / max(x["Total"].sum(), 1e-9)
    return (cs.clip(upper=0.999999) * n).apply(lambda v: int(v) + 1).reindex(df.index)


def opt_mix(df, n):
    x = df.sort_values("Total", ascending=False)
    b = []
    for i in range(len(x)):
        r, c = divmod(i, n)
        b.append(c + 1 if r % 2 == 0 else n - c)
    return pd.Series(b, index=x.index).reindex(df.index)


def build(months, label, out_xlsx, payload_path):
    mas = pd.read_csv(MASTER, dtype=str)
    off = vc.load_offtake(months)
    piv = off.pivot_table(index="sid", columns="file", values="NSV", aggfunc="sum").reindex(columns=months)
    units = off.groupby("sid")["Sales Qty"].sum().rename("Units")
    bcs = off.groupby("sid")["bc"].max().rename("bc")
    st = mas.set_index("Store Key")
    s = pd.concat([piv, units, bcs], axis=1).reset_index().rename(columns={"sid": "Store Key"})
    s = s.merge(mas, on="Store Key", how="left")
    s["Total"] = s[months].sum(axis=1, min_count=1)
    s["Chain"] = s["Chain Name"].fillna(s["Store Key"])
    s["Visit Status"] = s["Visit Status"].fillna("City not available")
    s["bc"] = s["bc"].fillna(False).astype(bool)

    cons = s[(s["Visit Status"] == "Considered") & s["Total"].notna()].copy()
    cons["Matched"] = cons["Visit City"]
    tot = cons.groupby("Matched")["Total"].sum().sort_values(ascending=False)
    top8 = list(tot.head(8).index)
    rank_pos = {c: i + 1 for i, c in enumerate(tot.index)}

    rows = []
    for city, g in cons.groupby("Matched"):
        n = 5 if city in top8 else 3
        g = g.copy()
        g["A"], g["B"], g["C"], g["n"] = opt_chain(g, n), opt_tier(g, n), opt_mix(g, n), n
        rows.append(g)
    beats = pd.concat(rows)

    # top articles per considered city
    art = pd.read_csv(vc.RAW / "offtake_store_article_Aug_26.csv", usecols=["EAN", "Description as per Fountain"], low_memory=False).dropna(subset=["EAN"]).drop_duplicates("EAN")
    name = dict(zip(art.EAN.astype("int64"), art["Description as per Fountain"]))
    city_of = cons.set_index("Store Key")["Matched"]
    dc = off[off.sid.isin(city_of.index)].copy()
    dc["city"] = dc.sid.map(city_of)
    dc["art"] = dc["EAN"].map(lambda e: name.get(int(e)) if pd.notna(e) else None).fillna(dc["EAN"].map(lambda e: f"EAN {int(e)}" if pd.notna(e) else "Unknown"))
    ta = dc.groupby(["city", "art"]).agg(NSV=("NSV", "sum"), Units=("Sales Qty", "sum"), Stores=("sid", "nunique")).reset_index()
    ta["rank"] = ta.groupby("city")["NSV"].rank(ascending=False, method="first")
    ta = ta[ta["rank"] <= 10].sort_values(["city", "rank"])

    nocity = s[s["Visit Status"] == "City not available"]
    nc = nocity.assign(State=nocity["State"].fillna("Not available")).groupby(["Chain", "State"])[months].sum(min_count=1).reset_index()
    nc["Total"] = nc[months].sum(axis=1, min_count=1)
    nc = nc.sort_values("Total", ascending=False)

    near = s[(s["Visit Status"] == "Near listed city") & s["Total"].notna()].copy()

    ly_raw = json.loads(LY_JSON.read_text(encoding="utf-8"))["by_chain"]
    lyu = {}
    for k, v in ly_raw.items():
        k = {"RATANDEEP": "RATNADEEP"}.get(k.upper(), k.upper())
        for mo, val in v.items():
            lyu.setdefault(k, {})
            lyu[k][mo] = lyu[k].get(mo, 0) + (val or 0)
    ty = off[~off.bc].groupby(["Chain", "file"])["NSV"].sum().unstack().reindex(columns=months)
    n_m = len(months)
    chain_yoy = {}
    for ch, row in ty.iterrows():
        a = sum(v for v in row if pd.notna(v))
        b = sum(lyu.get(ch, {}).get(LYM[i], 0) for i in range(n_m))
        chain_yoy[ch] = (a / b - 1) if b > 0 else None

    # master coverage
    on_master = set(mas.loc[mas["Source"] != "Offtake file, not in master", "Store Key"])
    sold = s[s["Total"].notna() & ~s["bc"]]
    cov = {"stores_sold": int(len(sold)), "in_master_file": int(sold["Store Key"].isin(on_master).sum()),
           "nsv_in_master_pct": float(sold.loc[sold["Store Key"].isin(on_master), "Total"].sum() / sold["Total"].sum() * 100)}

    # ---------------------------------------------------------------- Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    last_m = months[-1]
    lines = [
        ("City offtake, store list and tentative beat options", True),
        (f"Months in this file: {', '.join(m + ' 26' for m in months)}. Sep 2026 and last-year store sales are to be added (see NOT AVAILABLE below).", False),
        ("", False),
        ("HOW TO READ", True),
        ("1. Store_List: every store with a city. 'Consider city?' has three values: Considered (on your list), Near listed city (Navi Mumbai, Thane, Mohali, Panchkula, Ernakulam, Howrah: the nearest listed city is shown) and Not considered.", False),
        ("2. City_Summary: planned cities, region, stores, NSV by month and MoM. Near_Listed shows the six near-listed cities and the listed city each one rolls up to, so the team can decide whether to add them to that city's beats.", False),
        ("3. Beat_Options: three options for every planned store (A chain-wise, B sales tiers, C mixed). Beat_Summary shows what each beat looks like. Final beat and area: GT / EBO to fill.", False),
        ("4. City_Top_Articles: top 10 articles per planned city; stock availability and SOS columns are blank to fill.", False),
        ("5. Pack_Brand: for Facewash and Shampoo, which brands sell which pack size (Nielsen, IN URB MT), with Mamaearth's place in each pack.", False),
        ("6. Chain_NoCity: sales where no store city exists (chain / state level only). Chain_LastYear: chain level Apr-Sep 2025 vs this year. Master_Gaps: stores that sold but were not in the store master file.", False),
        ("", False),
        ("BASIS", True),
        ("NSV is Rs lakh as in the source files. Months come from the file names. Store city comes from the maintained store master (Store_City_Master.csv); where the master holds only a locality, or the store is new, the offtake file city is used (City Source column).", False),
        ("Reliance Brand Counter stores are listed as separate stores; city NSV is shown with and without them (the reported offtake basis excludes them).", False),
        (f"Store master coverage: {cov['in_master_file']:,} of {cov['stores_sold']:,} selling stores ({cov['nsv_in_master_pct']:.0f}% of NSV) are in the Jan-Jun 26 store list; the rest are new stores added from the offtake files (Master_Gaps).", False),
        ("Reliance (non-counter) is store-level only in Apr-Jun; in Jul-Aug it comes at state level, so its Jul-Aug sales sit in Chain_NoCity and city MoM for Jul-Aug is understated where Reliance stores exist.", False),
        ("", False),
        ("ASSUMPTIONS TO CONFIRM", True),
        (f"A. 'Top 8 metro' is not defined in the request. I took the 8 biggest listed cities by Apr-Aug NSV: {', '.join(top8)}.", False),
        ("B. Spelling: Cochin = Kochi, Trivandrum = Thiruvananthapuram, Trichy = Tiruchirappalli, Prayagraj = Allahabad, Vadodara = Baroda. Guwahati was listed twice and counted once.", False),
        ("C. Near-listed cities are NOT counted in the listed city's totals. They are shown as a third option so you can choose.", False),
        ("D. No address or pincode exists in the masters, so beats are groups of stores, not routes. Modern-trade stores only; GT and EBO outlets are not in these files.", False),
        ("", False),
        ("NOT AVAILABLE IN THE REPO (columns left blank, nothing estimated)", True),
        ("1. Last-year store-level sales (Apr-Sep 2025): needs the FY26 store x article offtake file (registry entry offtake_fy26_store_article, MISSING). The store master file lists stores only (no sales). Chain-level last year is in Chain_LastYear, and each store row shows its chain's YoY % for context (not a store figure).", False),
        ("2. Sep 2026 store-level offtake: not received yet; columns are ready.", False),
        ("3. Stock availability and SOS by article and store: no source in the repo.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(bold=b, size=12 if b else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 150

    # City summary
    ws = wb.create_sheet("City_Summary")
    mcols = [f"NSV {m}-26 (Rs L, excl BC)" for m in months]
    cols = ["Region", "City (as supplied)", "Beats planned", "Stores with sales"] + mcols + [f"MoM {m}" for m in months[1:]] + [
        "Total NSV excl BC", "Brand Counter NSV (separate)", "Near-listed NSV (not in totals)"] + [f"Last year {m}" for m in LYM] + ["Sep-26 (to add)", "Considered?", "Size rank"]
    grey = [c for c in cols if c.startswith("Last year") or c.startswith("Sep-26")]
    header(ws, 1, cols, grey)
    allc = [(r, c) for r, cs in vc.REGIONS.items() for c in cs]
    ncol_m = len(months)
    summary_rows = []
    for i, (reg, city) in enumerate(allc):
        r = 2 + i
        g = cons[cons.Matched == city]
        exb = g[~g.bc]
        nl = near[near["Nearest Listed City"] == city]
        vals = [reg, city, (5 if city in top8 else 3) if len(g) else 0, len(g)] + [fv(exb[m].sum(min_count=1)) for m in months]
        for j, v in enumerate(vals, 1):
            ws.cell(row=r, column=j, value=v)
        base = 5
        for k in range(1, ncol_m):
            cur, prv = L(base + k), L(base + k - 1)
            ws.cell(row=r, column=base + ncol_m + k - 1, value=f'=IFERROR({cur}{r}/{prv}{r}-1,"")').number_format = "0.0%"
        tcol = base + ncol_m + ncol_m - 1
        ws.cell(row=r, column=tcol, value=f"=SUM({L(base)}{r}:{L(base + ncol_m - 1)}{r})")
        ws.cell(row=r, column=tcol + 1, value=fv(g[g.bc]["Total"].sum()) if len(g) else None)
        ws.cell(row=r, column=tcol + 2, value=fv(nl["Total"].sum()) if len(nl) else None)
        for j in range(tcol + 3, tcol + 3 + len(LYM) + 1):
            ws.cell(row=r, column=j).fill = GFILL
        ws.cell(row=r, column=tcol + 3 + len(LYM) + 1, value="Considered" if len(g) else "Considered (no sales in files)")
        ws.cell(row=r, column=tcol + 3 + len(LYM) + 2, value=rank_pos.get(city))
        if city in top8:
            ws.cell(row=r, column=2).fill = YFILL
        summary_rows.append({"Region": reg, "City": city, "Beats": (5 if city in top8 else 3) if len(g) else 0, "Stores": len(g),
                             **{f"NSV_{m}": fv(exb[m].sum(min_count=1)) for m in months},
                             "NSV_total_exBC": fv(exb["Total"].sum()), "NSV_BC": fv(g[g.bc]["Total"].sum()),
                             "Near_listed_NSV": fv(nl["Total"].sum()), "Near_listed_stores": int(len(nl)), "Top8": city in top8})
    ws.cell(row=len(allc) + 3, column=2, value="Yellow = top 8 listed cities by NSV (5 beats). Grey = data not in the repo yet. Near-listed NSV is not added to the city totals.")
    widths(ws, [16, 18, 9, 10] + [13] * ncol_m + [9] * (ncol_m - 1) + [14, 16, 16] + [11] * 7 + [26, 8])

    # Store list
    ws = wb.create_sheet("Store_List")
    cols = ["Consider city?", "Visit city (supplied list)", "Nearest listed city", "Region", "Zone", "State", "City (final)", "City source", "City (store list)", "City (offtake file)",
            "Chain", "Store type", "Store code", "Store name"] + [f"NSV {m}-26 (Rs L)" for m in months] + [f"MoM {m}" for m in months[1:]] + [
        "Total NSV", "Units", "Chain YoY % (chain level, context)"] + [f"LY {m}" for m in LYM] + ["Sep-26 (to add)"]
    grey = [c for c in cols if c.startswith("LY") or c.startswith("Sep-26")]
    header(ws, 1, cols, grey)
    sl = s[s["Visit Status"].isin(["Considered", "Near listed city", "Not considered"]) & s["Total"].notna()].copy()
    sl["o"] = sl["Visit Status"].map({"Considered": 0, "Near listed city": 1, "Not considered": 2})
    sl = sl.sort_values(["o", "Visit Region", "Visit City", "Nearest Listed City", "City Final", "Total"], ascending=[True, True, True, True, True, False])
    nm = len(months)
    for i, (_, r) in enumerate(sl.iterrows(), 2):
        region = r["Visit Region"] if isinstance(r["Visit Region"], str) else vc.CITY_REGION.get(r["Nearest Listed City"])
        vals = [r["Visit Status"], r["Visit City"], r["Nearest Listed City"], region, r["Zone"], r["State"], r["City Final"], r["City Source"],
                r["City (store list)"], r["City (offtake file)"], r["Chain"].title() if isinstance(r["Chain"], str) else None, r["Store Type"], r["Site Code"], r["Store Name"]]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=fv(v))
        c0 = len(vals) + 1
        for k, m in enumerate(months):
            ws.cell(row=i, column=c0 + k, value=fv(r[m]))
        for k in range(1, nm):
            ws.cell(row=i, column=c0 + nm + k - 1, value=f'=IFERROR({L(c0 + k)}{i}/{L(c0 + k - 1)}{i}-1,"")').number_format = "0.0%"
        tc = c0 + nm + nm - 1
        ws.cell(row=i, column=tc, value=f"=SUM({L(c0)}{i}:{L(c0 + nm - 1)}{i})")
        ws.cell(row=i, column=tc + 1, value=fv(r["Units"]))
        y = chain_yoy.get(str(r["Chain"]).upper())
        ws.cell(row=i, column=tc + 2, value=fv(y)).number_format = "0%"
        for j in range(tc + 3, tc + 3 + len(LYM) + 1):
            ws.cell(row=i, column=j).fill = GFILL
        ws.cell(row=i, column=1).fill = GREEN if r["Visit Status"] == "Considered" else AMBER if r["Visit Status"] == "Near listed city" else RED
    ws.auto_filter.ref = f"A1:{L(len(cols))}{len(sl) + 1}"
    widths(ws, [16, 16, 14, 14, 10, 16, 18, 14, 18, 18, 16, 16, 12, 34] + [11] * nm + [8] * (nm - 1) + [12, 10, 14] + [10] * 7)

    # Near-listed sheet
    ws = wb.create_sheet("Near_Listed")
    cols = ["City (as in file)", "Nearest listed city", "State", "Stores with sales"] + [f"NSV {m}-26 (Rs L)" for m in months] + ["Total NSV", "Add to nearest city beats? (GT / EBO to decide)"]
    header(ws, 1, cols, ["Add to nearest city beats? (GT / EBO to decide)"])
    near["State"] = near["State"].fillna("Not available")
    ng = near.groupby(["City Final", "Nearest Listed City", "State"], dropna=False).agg(Stores=("Store Key", "count"), **{m: (m, lambda x: x.sum(min_count=1)) for m in months}, Total=("Total", "sum")).reset_index().sort_values("Total", ascending=False)
    for i, r in enumerate(ng.itertuples(index=False), 2):
        vals = [r[0], r[1], r[2], r[3]] + [fv(r[4 + k]) for k in range(nm)] + [fv(r[4 + nm]), None]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=len(cols)).fill = GFILL
    widths(ws, [28, 16, 18, 10] + [13] * nm + [12, 34])

    # Beat options
    ws = wb.create_sheet("Beat_Options")
    cols = ["Region", "City", "Beats in city", "Chain", "Store type", "Store code", "Store name", "Total NSV (Rs L)", "Option A: chain-wise beat", "Option B: sales-tier beat", "Option C: mixed beat",
            "Final beat (GT / EBO to fill)", "Area / route (GT / EBO to fill)"]
    header(ws, 1, cols, ["Final beat (GT / EBO to fill)", "Area / route (GT / EBO to fill)"])
    bs = beats.sort_values(["Visit Region", "Matched", "A", "Total"], ascending=[True, True, True, False])
    for i, (_, r) in enumerate(bs.iterrows(), 2):
        vals = [r["Visit Region"], r["Matched"], int(r["n"]), r["Chain"].title(), r["Store Type"], r["Site Code"], r["Store Name"], fv(r["Total"]), f"A{int(r['A'])}", f"B{int(r['B'])}", f"C{int(r['C'])}", None, None]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=fv(v))
        ws.cell(row=i, column=12).fill = GFILL
        ws.cell(row=i, column=13).fill = GFILL
    ws.auto_filter.ref = f"A1:{L(len(cols))}{len(bs) + 1}"
    widths(ws, [14, 14, 8, 16, 16, 12, 36, 14, 12, 12, 12, 18, 22])

    ws = wb.create_sheet("Beat_Summary")
    cols = ["Region", "City", "Option", "Beat", "Stores", "NSV (Rs L)", "% of city NSV", "Chains in beat", "Top 3 stores by NSV", "Brand Counter stores"]
    header(ws, 1, cols)
    r_ = 2
    for city, g in beats.groupby("Matched"):
        tot_c = g["Total"].sum()
        for opt, lab in (("A", "A chain-wise"), ("B", "B sales-tier"), ("C", "C mixed")):
            for b, x in g.groupby(opt):
                top = x.sort_values("Total", ascending=False).head(3)
                chains = ", ".join(f"{str(k).title()} ({v})" for k, v in x.Chain.value_counts().head(4).items())
                vals = [g["Visit Region"].iloc[0], city, lab, f"{opt}{int(b)}", len(x), fv(x["Total"].sum()), (x["Total"].sum() / tot_c) if tot_c else None, chains,
                        "; ".join(f"{(n or 'NA')} ({t:.1f})" for n, t in zip(top["Store Name"], top["Total"])), int(x.bc.sum())]
                for j, v in enumerate(vals, 1):
                    c = ws.cell(row=r_, column=j, value=fv(v))
                    if j == 7:
                        c.number_format = "0%"
                r_ += 1
    widths(ws, [14, 14, 14, 7, 8, 14, 10, 50, 80, 10])
    ws.auto_filter.ref = f"A1:J{r_ - 1}"

    ws = wb.create_sheet("City_Top_Articles")
    cols = ["City", "Rank", "Article", "NSV (Rs L)", "Units", "Stores selling", "Stock availability % (to fill)", "SOS % (to fill)", "Source for stock / SOS"]
    header(ws, 1, cols, cols[6:])
    for i, r in enumerate(ta.itertuples(index=False), 2):
        for j, v in enumerate([r.city, int(r.rank), r.art, fv(r.NSV), fv(r.Units), int(r.Stores)], 1):
            ws.cell(row=i, column=j, value=v)
        for j in (7, 8, 9):
            ws.cell(row=i, column=j).fill = GFILL
    widths(ws, [14, 6, 60, 14, 12, 10, 18, 12, 30])

    # Pack_Brand (Nielsen Brand x Basepack)
    ws = wb.create_sheet("Pack_Brand")
    pay = json.loads(Path(payload_path).read_text(encoding="utf-8"))
    r_ = 1
    for key, title in (("fw_pack_brand", "FACEWASH"), ("sh_pack_brand", "SHAMPOO")):
        pb = pay.get(key)
        if not pb:
            continue
        ws.cell(row=r_, column=1, value=f"{title}: brand share of each pack's value (%), Nielsen IN URB MT, {pb['month']}. Brands listed cover {pb.get('coverage_pct')}% of category value. Blank = not sold in the pack.").font = Font(bold=True)
        r_ += 1
        cols = [b["n"] for b in pb["brands"][:12]]
        if "Mamaearth" not in cols:
            cols.append("Mamaearth")
        hdr = ["Pack (ml)", "Category share %", "Category Rs Cr", "Brands selling", "Leader"] + cols
        for j, c in enumerate(hdr, 1):
            x = ws.cell(row=r_, column=j, value=c)
            x.font = HF
            x.fill = HFILL
        r_ += 1
        for p in pb["packs"]:
            ws.cell(row=r_, column=1, value=p["size"])
            ws.cell(row=r_, column=2, value=p["cat_share"])
            ws.cell(row=r_, column=3, value=p["cat_value"])
            ws.cell(row=r_, column=4, value=p["brands_selling"])
            ws.cell(row=r_, column=5, value=p["leader"])
            for j, c in enumerate(cols, 6):
                b = next((x for x in p["brands"] if x["n"] == c), None)
                if b and b["value"] > 0:
                    cell = ws.cell(row=r_, column=j, value=b["share_in_pack"])
                    if c == "Mamaearth":
                        cell.fill = GREEN
                    elif b["share_in_pack"] >= 30:
                        cell.fill = AMBER
                elif c == "Mamaearth":
                    ws.cell(row=r_, column=j, value="Not sold").fill = RED
            r_ += 1
        r_ += 1
        ws.cell(row=r_, column=1, value="Brand: packs sold and top packs").font = Font(bold=True)
        r_ += 1
        for b in pb["brands"][:15]:
            ws.cell(row=r_, column=1, value=b["n"])
            ws.cell(row=r_, column=2, value=b["packs"])
            ws.cell(row=r_, column=3, value=", ".join(b["top_packs"]))
            r_ += 1
        r_ += 2
    widths(ws, [14, 14, 14, 12, 18] + [13] * 14)

    ws = wb.create_sheet("Chain_NoCity")
    cols = ["Chain", "State"] + [f"NSV {m}-26 (Rs L)" for m in months] + ["Total", "Note"]
    header(ws, 1, cols)
    for i, r in enumerate(nc.itertuples(index=False), 2):
        vals = [str(r[0]).title(), r[1]] + [fv(r[2 + k]) for k in range(nm)] + [f"=SUM(C{i}:{L(2 + nm)}{i})", "Store city not in the sources; state / chain level only"]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
    widths(ws, [26, 20] + [13] * nm + [14, 50])

    ws = wb.create_sheet("Chain_LastYear")
    cols = ["Chain"] + [f"LY {x}" for x in LYM] + [f"TY {m}-26" for m in months] + ["TY Sep-26 (to add)"] + [f"YoY {m}" for m in months]
    header(ws, 1, cols, ["TY Sep-26 (to add)"])
    ty_up = ty.copy()
    ty_up.index = ty_up.index.str.upper()
    names = sorted(set(lyu) | set(ty_up.index), key=lambda c: -(ty_up.loc[c].sum() if c in ty_up.index else 0))
    names = [c for c in names if (c in lyu and any(lyu[c].values())) or (c in ty_up.index and ty_up.loc[c].notna().any())]
    for i, ch in enumerate(names, 2):
        ws.cell(row=i, column=1, value=ch.title())
        for k, mo in enumerate(LYM):
            ws.cell(row=i, column=2 + k, value=fv(lyu.get(ch, {}).get(mo)))
        for k, m in enumerate(months):
            ws.cell(row=i, column=8 + k, value=fv(ty_up.loc[ch, m]) if ch in ty_up.index else None)
        ws.cell(row=i, column=8 + nm).fill = GFILL
        for k in range(nm):
            ws.cell(row=i, column=9 + nm + k, value=f'=IFERROR({L(8 + k)}{i}/{L(2 + k)}{i}-1,"")').number_format = "0%"
    n = len(names) + 1
    ws.cell(row=n + 1, column=1, value="TOTAL").font = Font(bold=True)
    for c in range(2, 8 + nm):
        ws.cell(row=n + 1, column=c, value=f"=SUM({L(c)}2:{L(c)}{n})")
    ws.cell(row=n + 3, column=1, value="Chain level, all India, Rs lakh, offtake basis (Reliance Brand Counter excluded). Last year: data/raw_drops/_agg/offtake_fy26.json.")
    widths(ws, [26] + [11] * (7 + nm + nm))

    ws = wb.create_sheet("Master_Gaps")
    cols = ["Store key", "Chain", "State", "City (offtake file)", "Store name", "Total NSV (Rs L)", "Action"]
    header(ws, 1, cols)
    gaps = s[(s["Source"] == "Offtake file, not in master") & s["Total"].notna() & ~s["Store Key"].str.contains(r"\|NO-CITY", na=False)].sort_values("Total", ascending=False)
    for i, (_, r) in enumerate(gaps.iterrows(), 2):
        for j, v in enumerate([r["Store Key"], r["Chain"], r["State"], r["City (offtake file)"], r["Store Name"], fv(r["Total"]), "Add to the store master"], 1):
            ws.cell(row=i, column=j, value=fv(v))
    widths(ws, [34, 18, 16, 20, 34, 14, 22])
    wb.save(out_xlsx)

    summary = pd.DataFrame(summary_rows)
    near_rows = [{"city": r[0], "nearest_listed": r[1], "state": r[2], "stores": int(r[3]), "nsv_total": fv(r[4 + nm])} for r in ng.itertuples(index=False)]
    stats = {"stores_in_list": int(len(sl)), "gaps": int(len(gaps)), "near_stores": int(len(near)), "near_listed": near_rows,
             "status_counts": {k: int(v) for k, v in sl["Visit Status"].value_counts().items()}}
    return summary, top8, cov, stats


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--label", default="Aug26")
    ap.add_argument("--months", nargs="+", default=vc.MONTHS)
    ap.add_argument("--payload", type=Path, default=ROOT / "data" / "nielsen_aug26.json")
    a = ap.parse_args()
    summary, top8, cov, stats = build(a.months, a.label, a.out, a.payload)
    csv_path = ROOT / "data" / "nielsen" / f"Visit_City_Summary_{a.label}.csv"
    summary.to_csv(csv_path, index=False)
    (ROOT / "data" / "nielsen" / f"Visit_Cities_{a.label}.json").write_text(json.dumps({
        "months": [m + " 26" for m in a.months], "top8": top8, "coverage": cov, "stats": stats,
        "cities": json.loads(summary.to_json(orient="records")),
        "note": "Store master: PowerBI/SeedData/Masters/Store_City_Master.csv. NSV Rs lakh, offtake basis (Brand Counter separate)."}, indent=1) + "\n", encoding="utf-8")
    print("top 8:", top8)
    print("coverage:", cov, stats)
    print("wrote", a.out, "and", csv_path)


if __name__ == "__main__":
    main()
