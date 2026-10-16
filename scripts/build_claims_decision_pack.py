"""Decision pack for the open distributor-claim items in chain-wise CM2 (read-only: sets no decision, changes no data).

It lists, with numbers from the repo, what Finance / the business owner has to decide, and what each decision changes:
  1. the 7 driver aliases (PowerBI/SeedData/Mapping/ClaimDriver_Alias_Decisions_Q1_FY27.csv): secondary NSV behind each name, and
     the chain x month claim split now vs if all 7 were approved (Q1 totals never change, only the month timing);
  2. Ratnadeep (Q1 claim with no month) and Spar (claim with no NSV);
  3. the chains with sales but no cost loaded (no CM2 shown);
  4. chains whose returns are above sales.
An alias counts only when Owner_Decision = Approve with approver, date and reference (rule in build_distributor_claim_monthly_allocation.py).
After the decisions come back: fill the CSV, run build_distributor_claim_monthly_allocation.py, then build_dashboard_data.py --cm2-claims-patch.

  python scripts/build_claims_decision_pack.py [--out exports/CM2_Claims_Decision_Pack.xlsx]
"""
import argparse
import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import build_distributor_claim_monthly_allocation as A  # noqa: E402

# Confidence is the analyst's reading of the name evidence, shown so the approver can see where to look twice. Not a decision.
CONFIDENCE = {
    "MORE REATIL": ("High", "Typo of MORE RETAIL"),
    "MORE RETAIL LIMITED": ("High", "Legal name of More Retail"),
    "FRANK ROSE": ("High", "Typo of FRANK ROSS (the file also has FRANK ROSS)"),
    "RATNADEEP SUPERMARKET": ("High", "Store name of Ratnadeep; only Ratnadeep-like name in the file"),
    "ABRL": ("Check", "Aditya Birla Retail Ltd runs More, but ABRL may also cover other formats: confirm it is More only"),
    "ABRL -MORE": ("High", "Label itself says More"),
    "MRL": ("Check", "Initials only: confirm MRL is More Retail Ltd and not another MRL distributor"),
}
NAVY = "10254A"


def head(ws, r, n):
    for c in range(1, n + 1):
        x = ws.cell(row=r, column=c)
        x.font, x.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
        x.alignment = Alignment(wrap_text=True, vertical="center")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "exports" / "CM2_Claims_Decision_Pack.xlsx"))
    a = ap.parse_args()
    seed = list(csv.DictReader(A.SEED.open(encoding="utf-8-sig")))
    now = A.build()
    d, n, _ = A.load_driver(A.SECONDARY, approved=A.PENDING_ALIASES, pending=A.PENDING_ALIASES)     # as if all 7 approved
    new = A.allocate(seed, d, n, {})
    ev = defaultdict(lambda: defaultdict(float))
    for r in csv.DictReader(A.SECONDARY.open(encoding="utf-8-sig")):
        k = r["Chain"].strip()
        if k in A.PENDING_ALIASES:
            ev[k][r["Month_Label"]] += float(r["NSV_Lakh"] or 0)
    def by_cm(rows):
        t = defaultdict(float)
        for r in rows:
            t[(r["Chain"], r["Month"])] += float(r["Expense Amount (INR Lakh)"])
        return t
    bn, bw = by_cm(now), by_cm(new)
    dj = None
    s = (ROOT / "dashboard" / "data.js").read_text(encoding="utf-8")
    dj = json.loads(s[s.index("{"): s.rindex("}") + 1])
    cm = dj["cm2"]
    total_nsv = cm["total_nsv"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    lines = [
        ("CM2 distributor claims: open items and what each decision changes", True),
        (f"Claims: source Rs {cm['claims']['source_total']:.2f} L; applied Rs {cm['claims']['loaded']:.2f} L; not applied Rs {cm['claims']['unresolved_total']:.2f} L (Spar). Status ALLOCATED_PROVISIONAL: Finance (DA) 2026-10-01 A1/A2; formal approval reference pending.", False),
        ("1. Driver aliases (sheet Alias_decisions): 7 pending. Approving all 7 moves the monthly split for More Retail, Frankross and Ratnadeep. Q1 totals do not change. Ratnadeep's Rs 3.96 L gets a month.", False),
        ("2. Ratnadeep Rs 3.96 L has no month ONLY because RATNADEEP SUPERMARKET is not yet an approved alias. Spar Rs -0.02 L has no NSV in any source; see sheet Ratnadeep_Spar.", False),
        (f"3. {sum(1 for r in cm['by_chain'] if not r.get('expense'))} chains with sales but no cost loaded (sheet No_cost_chains): no direct DN and no claim in any file. Source gap, not a software error. The Direct DN input file covers 12 chains only.", False),
        ("4. Returns above sales (sheet Returns_above_sales): 4 chain-FY pairs, Rs -5.18 L, disclosed, not in the chain table.", False),
        ("What I cannot do: set any Owner_Decision, approve an alias, or invent cost for a chain. Fill the decision columns, send the CSV rows back, and the numbers update with two commands.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(bold=b, size=13 if b else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 22 if b else 46
    ws.column_dimensions["A"].width = 140

    wa = wb.create_sheet("Alias_decisions")
    hd = ["Secondary_Name", "Proposed_Chain", "Evidence (from file)", "Analyst read", "Confidence", "Secondary NSV Apr'26 (L)", "May'26 (L)", "Jun'26 (L)",
          "Owner_Decision (Approve / Reject)", "Approver", "Approval_Date", "Approval_Reference"]
    for j, h in enumerate(hd, 1):
        wa.cell(row=1, column=j, value=h)
    head(wa, 1, len(hd))
    for i, r in enumerate(csv.DictReader(A.ALIAS_TABLE.open(encoding="utf-8-sig")), 2):
        nm = r["Secondary_Name"]
        conf, why = CONFIDENCE.get(nm, ("", ""))
        vals = [nm, r["Proposed_Chain"], r["Evidence"], why, conf] + [round(ev[nm].get(m, 0.0), 2) for m, _ in A.MONTHS] + ["", "", "", ""]
        for j, v in enumerate(vals, 1):
            wa.cell(row=i, column=j, value=v)
    for c, w in zip("ABCDEFGHIJKL", (24, 16, 38, 46, 11, 14, 11, 11, 20, 14, 14, 18)):
        wa.column_dimensions[c].width = w

    wc = wb.create_sheet("Claims_before_after")
    hh = ["Chain", "Month", "Claim now (L)", "Claim if all 7 aliases approved (L)", "Change (L)"]
    for j, h in enumerate(hh, 1):
        wc.cell(row=1, column=j, value=h)
    head(wc, 1, 5)
    r = 2
    for k in sorted(set(bn) | set(bw)):
        if abs(bn.get(k, 0) - bw.get(k, 0)) < 0.00005:
            continue
        wc.cell(row=r, column=1, value=k[0]); wc.cell(row=r, column=2, value=k[1])
        wc.cell(row=r, column=3, value=round(bn.get(k, 0), 4)); wc.cell(row=r, column=4, value=round(bw.get(k, 0), 4))
        wc.cell(row=r, column=5, value=round(bw.get(k, 0) - bn.get(k, 0), 4))
        r += 1
    wc.cell(row=r + 1, column=1, value="Sum of the change column is 0: approving aliases only moves claims between months (and gives Ratnadeep its months). Q1 per chain is unchanged.")
    for c, w in zip("ABCDE", (22, 16, 16, 30, 14)):
        wc.column_dimensions[c].width = w

    wr = wb.create_sheet("Ratnadeep_Spar")
    rows = [
        ("Item", "Amount (L)", "Why open", "Options", "Needs"),
        ("Ratnadeep, no month", 3.9622, "Claim master names RATNADEEP RETAIL PVT LTD. The distributor secondary file names the chain RATNADEEP SUPERMARKET, which is a pending alias, so there is no Apr-Jun driver.", "A: approve the alias (claim splits by its Apr/May/Jun secondary: Rs 1.47 / 1.19 / 1.31 L). B: keep in the Q1 / FY total only.", "Alias decision (sheet Alias_decisions)"),
        ("Spar, not applied", -0.0205, "A credit of Rs 2,050 for a chain with no NSV in primary, offtake, secondary or CM2.", "A: Finance writes it off or holds it as non-material, stays disclosed. B: if Spar is a chain already in the master (the name may map to one, for example Max Hyper: I found no evidence either way), give the mapping and it will be applied.", "Finance: write off / hold / mapping"),
    ]
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = wr.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    head(wr, 1, 5)
    for c, w in zip("ABCDE", (22, 12, 60, 70, 34)):
        wr.column_dimensions[c].width = w

    wn = wb.create_sheet("No_cost_chains")
    for j, h in enumerate(["Chain", "NSV (L)", "% of total NSV", "Direct DN in PL_Expense_Input?", "Claim in master?", "Question for Finance"], 1):
        wn.cell(row=1, column=j, value=h)
    head(wn, 1, 6)
    dn = {(r.get("Chain") or "").strip() for r in csv.DictReader((ROOT / "PowerBI/SeedData/Masters/PL_Expense_Input.csv").open(encoding="utf-8-sig"))}
    clm = {r["Chain"] for r in seed}
    nocost = sorted((r for r in cm["by_chain"] if not r.get("expense")), key=lambda r: -r["nsv"])
    for i, r in enumerate(nocost, 2):
        vals = [r["name"], round(r["nsv"], 2), r["nsv"] / total_nsv, "Yes" if r["name"] in dn else "No", "Yes" if r["name"] in clm else "No",
                "Confirm: no direct DN and no distributor claim for this chain in FY27 (then CM2 can be shown as NSV less nil, flagged), or send the DN / claim file."]
        for j, v in enumerate(vals, 1):
            wn.cell(row=i, column=j, value=v)
        wn.cell(row=i, column=3).number_format = "0.00%"
    wn.cell(row=len(nocost) + 3, column=1, value="Total").font = Font(bold=True)
    wn.cell(row=len(nocost) + 3, column=2, value=f"=SUM(B2:B{len(nocost) + 1})").font = Font(bold=True)
    wn.cell(row=len(nocost) + 3, column=3, value=f"=SUM(C2:C{len(nocost) + 1})").number_format = "0.0%"
    for c, w in zip("ABCDEF", (30, 12, 14, 26, 18, 90)):
        wn.column_dimensions[c].width = w

    wt = wb.create_sheet("Returns_above_sales")
    for j, h in enumerate(["Chain", "FY", "NSV (L)", "Treatment"], 1):
        wt.cell(row=1, column=j, value=h)
    head(wt, 1, 4)
    for i, x in enumerate(cm["claims"].get("net_negative_nsv_chains_not_in_table", []), 2):
        for j, v in enumerate([x["chain"], x["fy"], x["nsv"], "Returns above sales: not a chain row (no CM2 ratio possible); in total NSV, disclosed"], 1):
            wt.cell(row=i, column=j, value=v)
    for c, w in zip("ABCD", (22, 10, 12, 80)):
        wt.column_dimensions[c].width = w
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}: {len(nocost)} no-cost chains, {sum(1 for _ in A.PENDING_ALIASES)} aliases")


if __name__ == "__main__":
    main()
