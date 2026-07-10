"""
Phase 3 (part B) — document ingest & offline summarisation.

Reads documents into plain text so the agent can summarise / answer questions
over reports that are locked in static formats. Text/markdown is handled with
the standard library (works everywhere, offline). PDF and XLSX use a *pluggable*
backend: whichever supported library is installed is used, and if none is
present a clear, actionable error is raised (never a silent wrong answer).

Summarisation prefers a real local model when one is supplied and reachable,
and otherwise falls back to a deterministic extractive summary — transparent
and offline, not a fabricated one.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Tuple

TEXT_EXTS = {".txt", ".md", ".markdown", ".rst", ".log", ""}


class DocumentError(Exception):
    """Unsupported or unreadable document."""


class DocumentDependencyError(DocumentError):
    """A backend library is required but not installed."""


@dataclass
class Document:
    path: str
    kind: str            # 'text' | 'pdf' | 'xlsx'
    text: str
    n_pages: Optional[int] = None
    meta: dict = field(default_factory=dict)


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------
def read_document(path) -> Document:
    p = Path(path)
    if not p.exists():
        raise DocumentError(f"File not found: {p}")
    ext = p.suffix.lower()
    if ext in TEXT_EXTS:
        return Document(str(p), "text", p.read_text(encoding="utf-8", errors="replace"))
    if ext == ".pdf":
        text, n = _read_pdf(p)
        return Document(str(p), "pdf", text, n_pages=n)
    if ext in (".xlsx", ".xlsm"):
        text = _read_xlsx(p)
        return Document(str(p), "xlsx", text)
    raise DocumentError(
        f"Unsupported document type {ext!r}. Supported: text/markdown, .pdf "
        f"(with a PDF backend), .xlsx (with openpyxl)."
    )


def _read_pdf(p: Path) -> Tuple[str, int]:
    """Extract text from a PDF using the first available backend."""
    # 1) pdfplumber
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(str(p)) as pdf:
            pages = [(pg.extract_text() or "") for pg in pdf.pages]
        return "\n".join(pages), len(pages)
    except ImportError:
        pass
    # 2) pypdf
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(str(p))
        pages = [(pg.extract_text() or "") for pg in reader.pages]
        return "\n".join(pages), len(pages)
    except ImportError:
        pass
    # 3) PyMuPDF
    try:
        import fitz  # type: ignore
        doc = fitz.open(str(p))
        pages = [pg.get_text() for pg in doc]
        return "\n".join(pages), len(pages)
    except ImportError:
        pass
    raise DocumentDependencyError(
        "No PDF backend installed. Install one (offline-capable): "
        "`pip install pdfplumber`  (or `pypdf`, or `pymupdf`)."
    )


def _read_xlsx(p: Path) -> str:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        raise DocumentDependencyError(
            "Reading .xlsx requires openpyxl: `pip install openpyxl`."
        )
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    chunks: List[str] = []
    for ws in wb.worksheets:
        chunks.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            chunks.append("\t".join("" if c is None else str(c) for c in row))
    return "\n".join(chunks)


# --------------------------------------------------------------------------
# Text stats & summarisation
# --------------------------------------------------------------------------
_STOPWORDS = set(
    "a an the and or but of to in on for with at by from as is are was were be "
    "been being this that these those it its his her their our your my we you they "
    "i he she them us not no do does did has have had will would can could should "
    "than then so such into over under about above below up down out if else while".split()
)


def split_sentences(text: str) -> List[str]:
    parts = re.split(r"(?<=[.!?])\s+", text.strip())
    return [s.strip() for s in parts if s.strip()]


def text_stats(text: str) -> dict:
    words = re.findall(r"\w+", text)
    return {
        "chars": len(text),
        "words": len(words),
        "unique_words": len({w.lower() for w in words}),
        "sentences": len(split_sentences(text)),
    }


def summarize_text(text: str, max_sentences: int = 5, provider=None) -> str:
    """Summarise text. Uses a reachable model if provided; otherwise a
    deterministic extractive summary (word-frequency sentence scoring)."""
    text = (text or "").strip()
    if not text:
        return ""

    # prefer a real, reachable model (e.g. local Ollama)
    if provider is not None and hasattr(provider, "complete"):
        try:
            if getattr(provider, "is_available", lambda: True)():
                out = provider.complete(
                    f"Summarise the following in at most {max_sentences} sentences:\n\n{text}",
                    system="You are a concise, faithful summariser.",
                )
                if out and out.strip():
                    return out.strip()
        except Exception:
            pass  # fall through to offline extractive summary

    sentences = split_sentences(text)
    if len(sentences) <= max_sentences:
        return " ".join(sentences)

    freq: dict = {}
    for w in re.findall(r"\w+", text.lower()):
        if w in _STOPWORDS or len(w) <= 2:
            continue
        freq[w] = freq.get(w, 0) + 1

    def score(sentence: str) -> float:
        words = [w for w in re.findall(r"\w+", sentence.lower()) if w not in _STOPWORDS]
        if not words:
            return 0.0
        return sum(freq.get(w, 0) for w in words) / len(words)

    ranked = sorted(range(len(sentences)), key=lambda i: score(sentences[i]), reverse=True)
    chosen = sorted(ranked[:max_sentences])  # keep original order
    return " ".join(sentences[i] for i in chosen)
