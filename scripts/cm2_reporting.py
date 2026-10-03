"""Signed Primary source-tax QC, with all monetary values in INR lakh.

Callers must validate invoice NSV semantics and explicitly set
``frame.attrs['invoice_nsv_tax_exclusive'] = True`` before a rate is shown.
QC does not certify source units, source-to-allocation reconciliation, or
Finance approval. MRP is recorded independently and is never reconstructed.
"""
from pathlib import Path

import numpy as np
import pandas as pd

# A 0.05 percentage point tolerance accommodates recorded rounding. Nonzero
# tax on less than Rs 1 NSV or Rs 0.10 tax is reviewed even if its ratio fits.
# These are QC thresholds, not authority to adjust or discard source amounts.
RATE_TOLERANCE_PCT = 0.05
MIN_ABS_NSV_LAKH = 1.0 / 100_000
MIN_ABS_TAX_LAKH = 0.10 / 100_000
REQUIRED_COLUMNS = ("_FY", "_M", "_Chain", "_Brand", "_category", "_EAN No.",
                    "_NSV", "_TaxLOC", "_MRP", "_Chan")


def build_tax_basis(primary: pd.DataFrame) -> pd.DataFrame:
    """Copy allocated article rows, preserve signed source fields, add row QC.

    Missing tax is distinct from recorded zero. Missing/nonfinite base or MRP,
    zero base, inconsistent signs, and immaterial nonzero invoices are reviewed.
    The function neither changes amounts nor asserts tax exclusivity for callers.
    """
    missing = sorted(set(REQUIRED_COLUMNS) - set(primary.columns))
    if missing:
        raise ValueError(f"Missing Primary columns: {', '.join(missing)}")
    rows = primary.copy()
    nsv, tax, mrp = (pd.to_numeric(rows[c], errors="coerce")
                     for c in ("_NSV", "_TaxLOC", "_MRP"))
    status = pd.Series("REVIEW_RATE_OR_BASE", index=rows.index, dtype="object")
    compatible = (np.isfinite(nsv) & np.isfinite(tax) & np.isfinite(mrp)
                  & nsv.ne(0) & (np.sign(mrp) == np.sign(nsv)))
    status.loc[compatible & tax.eq(0)] = "ZERO_TAX"
    material = (compatible & nsv.abs().ge(MIN_ABS_NSV_LAKH)
                & tax.abs().ge(MIN_ABS_TAX_LAKH) & (np.sign(tax) == np.sign(nsv)))
    rate = 100 * tax / nsv.where(nsv.ne(0))
    for expected in (5, 18):
        status.loc[material & (rate - expected).abs().le(RATE_TOLERANCE_PCT)] = f"VALID_{expected}_PERCENT"
    status.loc[tax.isna()] = "MISSING_TAX"
    rows["gst_qc_status"] = status
    return rows


def tax_summary(rows: pd.DataFrame) -> dict:
    """Summarize a selected signed population, withholding unvalidated rates.

    ``reviewed_amount`` is absolute reviewed NSV exposure (INR lakh), preventing
    credits from cancelling exceptions; ``reviewed_nsv``/``reviewed_tax`` remain
    signed. Missing tax sums stay None if no numeric tax exists. A population
    with missing/review rows exposes its recorded totals but no effective rate.
    """
    if "gst_qc_status" not in rows:
        raise ValueError("Call build_tax_basis before tax_summary")
    amounts = {key: pd.to_numeric(rows[col], errors="coerce")
               for key, col in {"nsv": "_NSV", "tax": "_TaxLOC", "mrp": "_MRP"}.items()}
    def total(values):
        value = values.sum(min_count=1)
        return float(value) if np.isfinite(value) else None
    result = {key: total(values) for key, values in amounts.items()}
    counts = {str(k): int(v) for k, v in rows.gst_qc_status.value_counts().items()}
    statuses = set(counts)
    if "MISSING_TAX" in statuses:
        aggregate = "MISSING_TAX"
    elif "REVIEW_RATE_OR_BASE" in statuses:
        aggregate = "REVIEW_RATE_OR_BASE"
    elif len(statuses) > 1:
        aggregate = "MIXED_RATE"
    else:
        aggregate = next(iter(statuses), "NOT_AVAILABLE")
    reviewed = rows.gst_qc_status.eq("REVIEW_RATE_OR_BASE")
    missing = rows.gst_qc_status.eq("MISSING_TAX")
    validated = rows.attrs.get("invoice_nsv_tax_exclusive") is True
    rate_available = (validated and bool(statuses)
                      and statuses <= {"VALID_5_PERCENT", "VALID_18_PERCENT", "ZERO_TAX"}
                      and result["nsv"] is not None and result["nsv"] != 0
                      and result["tax"] is not None)
    result.update(
        effective_gst_pct=100 * result["tax"] / result["nsv"] if rate_available else None,
        invoice_nsv_tax_exclusive=validated,
        gst_qc_status=aggregate,
        status_counts=counts,
        reviewed_amount=float(amounts["nsv"][reviewed].abs().sum()),
        reviewed_nsv=total(amounts["nsv"][reviewed]),
        reviewed_tax=total(amounts["tax"][reviewed]),
        missing_tax_nsv=total(amounts["nsv"][missing]),
    )
    return result

# Input headers are a source contract, not a fuzzy column-name match.
_WORKBOOK_HEADERS = {
    'Chain_Master': ('P&L chain', 'Format', 'TOT/MD %', 'Source of TOT/MD'),
    'Chain_Alias': ('Source system', 'Chain name AS IT APPEARS in the source', 'P&L chain (locked name from your Provision sheet)', 'Note'),
    'Data_Primary': ('Month', 'P&L chain (auto from Chain_Alias)', 'Brand', 'Primary NSV (Rs Lakh)', 'Chain name in source', 'Source'),
    'Data_Provision': ('Month', 'P&L chain', 'Head', 'Amount (Rs Lakh)', 'Type', 'Your sheet cell', 'In Chain_Master?'),
    'Data_ActualClaims': ('Month', 'P&L chain (auto from Chain_Alias)', 'Expense head', 'Actual claim - DN (Rs Lakh)', 'Chain name in source', 'Month status', 'Source'),
}
_CUT_COLUMNS = ['month', 'chain', 'head', 'amount_lakh', 'source_row',
                'source_status', 'mapping_status', 'source_chain',
                'source_event_id', 'source_type']


def load_workbook_cuts(path: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Read provision Claim, provision BA, and DN without Excel formula caches.

    Months are calendar YYYY-MM. Lineage is local: sheet row and a source-event
    ID derived from provision month/cell or DN month/source-chain/head/source.
    Duplicate event keys are rejected; DN is an aggregated register and has no
    invoice ID. ``source_status`` describes evidence, not Finance approval.
    Unmapped rows retain their source name and signed value. Each cut's attrs
    holds ``chain_aliases`` and normalized ``workbook_primary`` (rounded source
    NSV) for independent reconciliation; these attrs are never public payloads.
    """
    import hashlib
    import json
    from datetime import datetime
    import openpyxl

    def text(value, location):
        if not isinstance(value, str) or not value.strip() or value.startswith('='):
            raise ValueError(f'Invalid source text at {location}')
        return value.strip()

    def month(value, location):
        try:
            return datetime.strptime(text(value, location), '%b-%y').strftime('%Y-%m')
        except ValueError as exc:
            raise ValueError(f'Invalid month at {location}: {value!r}') from exc

    def amount(value, location):
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not np.isfinite(value):
            raise ValueError(f'Invalid amount at {location}: {value!r}')
        return float(value)

    w = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = {}
        for name, expected in _WORKBOOK_HEADERS.items():
            if name not in w.sheetnames:
                raise ValueError(f'Missing workbook sheet: {name}')
            values = iter(w[name].values)
            if tuple(next(values, ())) != expected:
                raise ValueError(f'Unexpected column contract for {name}')
            sheets[name] = [(row_number, row) for row_number, row in enumerate(values, 2)
                            if any(v is not None for v in row)]
        locked = {text(r[0], f'Chain_Master!A{n}') for n, r in sheets['Chain_Master']}
        aliases = {}
        for n, r in sheets['Chain_Alias']:
            source, target = text(r[1], f'Chain_Alias!B{n}'), text(r[2], f'Chain_Alias!C{n}')
            if target not in locked:
                raise ValueError(f'Alias target outside Chain_Master at Chain_Alias!C{n}')
            if source in aliases and aliases[source] != target:
                raise ValueError(f'Conflicting alias at Chain_Alias!B{n}: {source}')
            aliases[source] = target
        primary = []
        for n, r in sheets['Data_Primary']:
            source = text(r[4], f'Data_Primary!E{n}')
            primary.append(dict(month=month(r[0], f'Data_Primary!A{n}'),
                                chain=aliases.get(source, 'UNMAPPED'),
                                brand=text(r[2], f'Data_Primary!C{n}'),
                                nsv=amount(r[3], f'Data_Primary!D{n}'),
                                source_chain=source, source_row=n))
        primary = pd.DataFrame(primary, columns=['month', 'chain', 'brand', 'nsv', 'source_chain', 'source_row'])
        cuts = {'PROVISION_CLAIM': [], 'PROVISION_BA': [], 'DIRECT_DN_ACTUAL': []}
        seen = {}
        for sheet in ('Data_Provision', 'Data_ActualClaims'):
            for n, r in sheets[sheet]:
                period = month(r[0], f'{sheet}!A{n}')
                head = text(r[2], f'{sheet}!C{n}')
                value = amount(r[3], f'{sheet}!D{n}')
                if sheet == 'Data_Provision':
                    if r[4] not in ('Claim', 'BA'):
                        raise ValueError(f'Unsupported Type at {sheet}!E{n}: {r[4]!r}')
                    source = text(r[1], f'{sheet}!B{n}')
                    chain = source if source in locked else 'UNMAPPED'
                    kind = 'PROVISION_CLAIM' if r[4] == 'Claim' else 'PROVISION_BA'
                    status = 'PROVISION'
                    event = (sheet, period, text(r[5], f'{sheet}!F{n}'))
                else:
                    status_map = {'Full': 'RECORDED_DN', 'PARTIAL MONTH': 'PARTIAL_REGISTER'}
                    if r[5] not in status_map:
                        raise ValueError(f'Unsupported status at {sheet}!F{n}: {r[5]!r}')
                    status = status_map[r[5]]
                    source = text(r[4], f'{sheet}!E{n}')
                    chain = aliases.get(source, 'UNMAPPED')
                    kind = 'DIRECT_DN_ACTUAL'
                    event = (sheet, period, source, head, text(r[6], f'{sheet}!G{n}'))
                if event in seen:
                    raise ValueError(f'Duplicate source-event ID at {sheet} row {n}; first row {seen[event]}')
                seen[event] = n
                event_id = hashlib.sha256(json.dumps(event, ensure_ascii=False).encode()).hexdigest()
                cuts[kind].append(dict(month=period, chain=chain, head=head,
                                       amount_lakh=value, source_row=n, source_status=status,
                                       mapping_status='UNMAPPED' if chain == 'UNMAPPED' else 'MAPPED',
                                       source_chain=source, source_event_id=event_id, source_type=kind))
        frames = tuple(pd.DataFrame(cuts[kind], columns=_CUT_COLUMNS)
                       for kind in ('PROVISION_CLAIM', 'PROVISION_BA', 'DIRECT_DN_ACTUAL'))
        for frame in frames:
            frame.attrs.update(chain_aliases=aliases.copy(), workbook_primary=primary.copy(),
                               source_version=hashlib.sha256(Path(path).read_bytes()).hexdigest())
        return frames
    finally:
        w.close()


def claim_coverage(primary: pd.DataFrame, claims: pd.DataFrame) -> dict:
    """Reconcile a selected normalized ``month, chain, nsv`` sales population.

    Use calendar YYYY-MM and canonical chains on both sides. Signed nonzero
    chain-month NSV is coverable, including negative-only groups. Net-zero and
    absent sales leave claims unallocated. Unmapped amounts have their own
    bucket, even when unmapped Primary exists. No amount or percentage is rounded.
    """
    for frame, required, name in ((primary, {'month', 'chain', 'nsv'}, 'Primary'),
                                  (claims, {'month', 'chain', 'amount_lakh', 'source_row'}, 'claims')):
        missing = required - set(frame.columns)
        if missing:
            raise ValueError(f'Missing {name} columns: {sorted(missing)}')
        numeric = pd.to_numeric(frame['nsv' if name == 'Primary' else 'amount_lakh'], errors='coerce')
        if not np.isfinite(numeric).all():
            raise ValueError(f'Nonfinite {name} amount')
        if frame[['month', 'chain']].isna().any().any():
            raise ValueError(f'Missing {name} coverage key')
    sales = primary.assign(nsv=pd.to_numeric(primary.nsv)).groupby(['month', 'chain'], sort=True).nsv.sum()
    mapped = claims.chain.ne('UNMAPPED')
    if 'mapping_status' in claims:
        mapped &= claims.mapping_status.ne('UNMAPPED')
    claim_keys = set(map(tuple, claims.loc[mapped, ['month', 'chain']].values))
    available = {key for key, value in sales.items() if value != 0 and key[1] != 'UNMAPPED'}
    covered = claim_keys & available
    no_sales = claim_keys - available
    unallocated, allocated, no_sales_amount, unmapped_amount = [], 0.0, 0.0, 0.0
    for is_mapped, (_, row) in zip(mapped, claims.iterrows()):
        key = (row['month'], row['chain'])
        value = float(row['amount_lakh'])
        if not is_mapped:
            status = 'UNMAPPED'
            unmapped_amount += value
        elif key in no_sales:
            status = 'UNALLOCATED_NO_SALES'
            no_sales_amount += value
        else:
            allocated += value
            continue
        unallocated.append(dict(month=key[0], chain=key[1], amount_lakh=value,
                                source_row=int(row['source_row']), status=status))
    return dict(covered_nsv=float(sum(v for k, v in sales.items() if k in covered)),
                uncovered_nsv=float(sum(v for k, v in sales.items() if k not in covered)),
                covered_keys=sorted(covered), no_sales_keys=sorted(no_sales),
                no_sales_claim_amount=no_sales_amount, unmapped_claim_amount=unmapped_amount,
                allocated_claim_amount=allocated,
                source_total=float(pd.to_numeric(claims.amount_lakh).sum()), unallocated=unallocated)


CM2_SCHEMA_VERSION = '2'
CM2_EXPORT_FIELDS = (
    'schema_version', 'view', 'record_type', 'fy', 'month', 'chain', 'brand',
    'channel', 'nsv', 'tax', 'mrp', 'effective_gst_pct', 'gst_qc_status',
    'effective_gst_pct_status', 'claim_amount', 'ba_amount', 'covered_nsv',
    'coverage_pct', 'coverage_pct_status', 'source_status', 'status',
    'cm2_value', 'cm2_pct', 'cm2_pct_status', 'allocation_basis',
    'category_status', 'cogs_basis_status', 'approval_status',
    'assumption_version', 'included_claim_heads', 'source_version',
    'timing_status', 'tax_row_count', 'gst_valid_5_count', 'gst_valid_18_count',
    'gst_zero_tax_count', 'gst_review_count', 'gst_missing_tax_count',
    'reviewed_amount', 'reviewed_nsv', 'reviewed_tax', 'missing_tax_nsv',
)


def build_cm2_views(primary, provision_claims, provision_ba, direct_dn, *, cogs_basis_status):
    """Public aggregate contract. Brand attribution is a signed NSV-share estimate.

    Finance has not approved a COGS base: even a caller-supplied approval label
    cannot enable margin arithmetic without an implemented governed formula.
    None cuts mean absent evidence; empty cuts mean loaded, recorded zero rows.
    """
    import calendar
    for cut, expected in ((provision_claims, 'PROVISION_CLAIM'),
                          (provision_ba, 'PROVISION_BA'), (direct_dn, 'DIRECT_DN_ACTUAL')):
        if cut is not None and 'source_type' in cut and not cut.source_type.eq(expected).all():
            raise ValueError(f'Expected exclusively {expected} source rows')
    basis = build_tax_basis(primary)
    months = {m: i for names in (calendar.month_abbr, calendar.month_name)
              for i, m in enumerate(names) if m}
    months['Sept'] = 9  # canonical dashboard label
    month_number = basis['_M'].map(months)
    if month_number.isna().any() or not basis['_FY'].astype(str).str.fullmatch(r'FY\d{2}').all():
        raise ValueError('Invalid Primary FY/month key')
    years = 2000 + basis['_FY'].astype(str).str[2:].astype(int) - month_number.ge(4).astype(int)
    basis['month'] = years.astype(str) + '-' + month_number.astype(int).astype(str).str.zfill(2)
    aliases = next((c.attrs.get('chain_aliases', {}) for c in
                    (provision_claims, direct_dn) if c is not None), {})
    basis['chain'] = basis['_Chain'].map(lambda s: aliases.get(s, s))
    mt = basis.loc[basis['_Chan'].eq('MT')].copy()
    sales = mt[['month', 'chain', '_NSV']].rename(columns={'_NSV': 'nsv'})
    empty = pd.DataFrame(columns=['month', 'chain', 'amount_lakh', 'source_row'])
    views = {'contract_version': CM2_SCHEMA_VERSION, 'unit': 'INR Lakh',
             'channel_basis': 'MT', 'all_channel_source': tax_summary(basis),
             'assumptions': {'cogs_basis_status': cogs_basis_status,
                 'formula_status': 'NOT_AVAILABLE', 'approval_status': 'PENDING_FINANCE',
                 'cogs_rate': .16, 'logistics_rate': .03,
                 'rates_applied': False, 'version': 'unapproved-workbook-assumptions-v1'}}
    coverage = {}
    for name, claims, ba in [('provision', provision_claims, provision_ba),
                             ('recorded_dn', direct_dn, None)]:
        cut = claims if claims is not None else empty
        cov = claim_coverage(sales, cut)
        coverage[name] = cov
        totals = sales.groupby(['month', 'chain']).nsv.sum()
        amounts = cut.groupby(['month', 'chain']).amount_lakh.sum()
        ba_cut = ba if ba is not None else empty
        ba_cov = claim_coverage(sales, ba_cut)
        ba_totals = ba_cut.groupby(['month', 'chain']).amount_lakh.sum()
        rows = []
        for (fy, month, chain, brand), group in mt.groupby(
                ['_FY', 'month', 'chain', '_Brand'], dropna=False, sort=True):
            key = (month, chain)
            covered = key in cov['covered_keys']
            source_rows = cut.loc[cut.month.eq(month) & cut.chain.eq(chain)]
            statuses = set(source_rows.get('source_status', []))
            source_status = ('PARTIAL_REGISTER' if 'PARTIAL_REGISTER' in statuses
                             else ('PROVISION' if name == 'provision' else 'RECORDED_DN')
                             if covered else 'NOT_AVAILABLE')
            raw = tax_summary(group)
            share = raw['nsv'] / float(totals[key]) if totals[key] != 0 else None
            rows.append(dict(fy=fy, month=month, chain=chain, brand=brand,
                channel='MT', nsv=raw['nsv'], tax=raw['tax'], mrp=raw['mrp'],
                effective_gst_pct=raw['effective_gst_pct'], gst_qc_status=raw['gst_qc_status'],
                effective_gst_pct_status=raw['gst_qc_status'] if raw['effective_gst_pct'] is not None else 'NOT_AVAILABLE',
                claim_amount=float(amounts[key])*share if covered else None,
                ba_amount=float(ba_totals[key])*share if key in ba_cov['covered_keys'] else None,
                covered_nsv=raw['nsv'] if covered else 0.,
                coverage_pct=100. if covered else 0., coverage_pct_status='AVAILABLE',
                source_status=source_status, status='NOT_AVAILABLE',
                cm2_value=None, cm2_pct=None, cm2_pct_status='NOT_AVAILABLE',
                allocation_basis='NSV_SHARE_ESTIMATE', category_status='UNSUPPORTED_SCOPE'))
            rows[-1].update(cogs_basis_status=cogs_basis_status,
                approval_status='PENDING_FINANCE', assumption_version='unapproved-workbook-assumptions-v1',
                included_claim_heads='|'.join(sorted(set(source_rows.get('head', [])))),
                source_version=claims.attrs.get('source_version', 'UNVERSIONED_INPUT') if claims is not None else 'ABSENT',
                timing_status='PARTIAL_REGISTER' if source_status == 'PARTIAL_REGISTER' else 'NOT_CERTIFIED')
            counts = raw['status_counts']
            rows[-1].update(schema_version=CM2_SCHEMA_VERSION, record_type='SALES',
                tax_row_count=len(group), gst_valid_5_count=counts.get('VALID_5_PERCENT', 0),
                gst_valid_18_count=counts.get('VALID_18_PERCENT', 0),
                gst_zero_tax_count=counts.get('ZERO_TAX', 0),
                gst_review_count=counts.get('REVIEW_RATE_OR_BASE', 0),
                gst_missing_tax_count=counts.get('MISSING_TAX', 0),
                **{k: raw[k] for k in ('reviewed_amount', 'reviewed_nsv', 'reviewed_tax', 'missing_tax_nsv')})
        # Local row lineage is deliberately excluded from the public contract.
        def public_unallocated(c):
            totals = {}
            for row in c['unallocated']:
                key = (row['month'], row['chain'], row['status'])
                totals[key] = totals.get(key, 0.) + row['amount_lakh']
            return [dict(month=m, chain=ch, status=s, amount_lakh=amount)
                    for (m, ch, s), amount in sorted(totals.items())]
        public_cov = {k:v for k,v in cov.items() if k != 'unallocated'}
        views[name] = dict(status='NOT_AVAILABLE', source_available=claims is not None,
            rows=rows, coverage=public_cov, unallocated=public_unallocated(cov),
            ba_unallocated=public_unallocated(ba_cov), ba_source_total=ba_cov['source_total'])
        exception_rows = []
        for bucket, source, kind, measure in (
            ('unallocated', claims, 'UNALLOCATED_CLAIM', 'claim_amount'),
            ('ba_unallocated', ba, 'UNALLOCATED_BA', 'ba_amount')):
            for entry in views[name][bucket]:
                year, month_number = map(int, entry['month'].split('-'))
                evidence = source.loc[source.month.eq(entry['month']) & source.chain.eq(entry['chain'])]
                partial = 'PARTIAL_REGISTER' in set(evidence.get('source_status', []))
                row = {field: None for field in CM2_EXPORT_FIELDS if field != 'view'}
                row.update(schema_version=CM2_SCHEMA_VERSION, record_type=kind,
                    fy=f'FY{(year + (month_number >= 4)) % 100:02d}',
                    month=entry['month'], chain=entry['chain'], channel='MT',
                    status=entry['status'], source_status='PARTIAL_REGISTER' if partial else 'PROVISION' if name == 'provision' else 'RECORDED_DN',
                    covered_nsv=0., coverage_pct_status='NOT_AVAILABLE',
                    effective_gst_pct_status='NOT_AVAILABLE', gst_qc_status='NOT_AVAILABLE',
                    cm2_pct_status='NOT_AVAILABLE', cogs_basis_status=cogs_basis_status,
                    approval_status='PENDING_FINANCE', allocation_basis='UNALLOCATED',
                    category_status='UNSUPPORTED_SCOPE', timing_status='PARTIAL_REGISTER' if partial else 'NOT_CERTIFIED',
                    assumption_version='unapproved-workbook-assumptions-v1',
                    included_claim_heads='|'.join(sorted(set(evidence.get('head', [])))),
                    source_version=source.attrs.get('source_version', 'UNVERSIONED_INPUT'),
                    **{measure:entry['amount_lakh']})
                for field in ('tax_row_count','gst_valid_5_count','gst_valid_18_count',
                              'gst_zero_tax_count','gst_review_count','gst_missing_tax_count'):
                    row[field] = 0
                exception_rows.append(row)
        views[name]['exception_rows'] = exception_rows
    matched = set(coverage['provision']['covered_keys']) & set(coverage['recorded_dn']['covered_keys'])
    denominator = float(sales.loc[[tuple(r) in matched for r in sales[['month','chain']].values], 'nsv'].sum())
    def matched_amount(c):
        if c is None:
            return None
        return float(c.loc[[tuple(r) in matched for r in c[['month','chain']].values], 'amount_lakh'].sum())
    pc, dn = matched_amount(provision_claims), matched_amount(direct_dn)
    gap = dn-pc if pc is not None and dn is not None else None
    views['matched_bridge'] = dict(key_count=len(matched), keys=sorted(matched), nsv=denominator,
        provision_claims=pc, recorded_dn_claims=dn, claim_difference=gap,
        claim_gap_pct=100*gap/denominator if gap is not None and denominator else None,
        claim_gap_pct_status='OBSERVED_CLAIM_GAP' if denominator and gap is not None else 'NOT_AVAILABLE',
        status='NOT_COMPARABLE', reason='Plan basis and recorded register timing differ')
    workbook = next((c.attrs.get('workbook_primary') for c in (provision_claims, direct_dn)
                     if c is not None and 'workbook_primary' in c.attrs), None)
    # Workbook groups minor brands into "Other brands". Compare monetary
    # populations on canonical chain/month, retaining every signed raw brand.
    if workbook is not None:
        keys = set(map(tuple, workbook[['month','chain']].values))
        scoped = mt.loc[[tuple(r) in keys for r in mt[['month','chain']].values]]
        raw_nsv, rounded_nsv = float(scoped._NSV.sum()), float(workbook.nsv.sum())
    else:
        raw_nsv, rounded_nsv = float(mt._NSV.sum()), None
    views['rounding_bridge'] = dict(raw_nsv=raw_nsv, workbook_nsv=rounded_nsv,
        raw_minus_workbook_nsv=raw_nsv-rounded_nsv if rounded_nsv is not None else None,
        scope='WORKBOOK_CHAIN_MONTH_KEYS' if workbook is not None else 'ALL_MT')
    return views


def write_cm2_powerbi_export(views: dict, path: Path) -> None:
    """Versioned sales and aggregate exceptions; never lose claim/BA amounts.

    Consumers must filter record_type for sales/GST denominators. Exception
    rows carry no brand, NSV or tax; signed claim and BA totals are conserved.
    The full fixed schema is written even when the population is empty.
    """
    import csv
    rows = [dict(view=name, **r) for name in ('provision', 'recorded_dn')
            for r in views[name]['rows'] + views[name]['exception_rows']]
    def write(stream):
        writer = csv.DictWriter(stream, fieldnames=CM2_EXPORT_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    if hasattr(path, 'write'):
        write(path)
    else:
        with Path(path).open('w', newline='', encoding='utf-8') as stream:
            write(stream)
