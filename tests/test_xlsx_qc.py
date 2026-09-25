"""
tests/test_xlsx_qc.py -- read-only Excel QC scanner (scripts/xlsx_qc.py).

Workbooks are built in tmp_path with openpyxl (already a repo dependency);
no real business file is used. Covers the recovery acceptance list: clean =>
PASS, #REF! detected, duplicate key detected, case and spacing variants
detected, corrupt/unsupported => named error, no network, no mutation.
"""
import hashlib
import socket
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import xlsx_qc  # noqa: E402


def _book(path, rows, sheet="Data", extra=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    if extra:
        extra(wb)
    wb.save(path)
    return path


def _checks(path, **kw):
    return [f.check for f in xlsx_qc.scan(xlsx_qc.read_workbook(path), **kw)]


CLEAN = [["Store Code", "Zone", "NSV"], ["S1", "West", 10], ["S2", "East", 20], ["S3", "North", 5]]


def test_clean_workbook_passes(tmp_path):
    p = _book(tmp_path / "clean.xlsx", CLEAN)
    assert _checks(p, keys=[("Data", ["Store Code"])]) == []
    assert xlsx_qc.main([str(p), "--key", "Data:Store Code"]) == 0


def test_ref_error_value_detected(tmp_path):
    p = _book(tmp_path / "ref.xlsx", CLEAN + [["S4", "#REF!", 1]])
    found = xlsx_qc.scan(xlsx_qc.read_workbook(p))
    err = [f for f in found if f.check == "ERROR_VALUE"]
    assert err and err[0].status == "FAIL" and err[0].cells == ["B5"] and "#REF!" in err[0].detail
    assert xlsx_qc.main([str(p)]) == 1


def test_ref_inside_formula_detected(tmp_path):
    p = _book(tmp_path / "reff.xlsx", CLEAN + [["S4", "South", "=#REF!+C2"]])
    assert "REF_IN_FORMULA" in _checks(p)


def test_duplicate_key_detected(tmp_path):
    p = _book(tmp_path / "dup.xlsx", CLEAN + [["S2", "East", 7]])
    f = [x for x in xlsx_qc.scan(xlsx_qc.read_workbook(p), keys=[("Data", ["Store Code"])])
         if x.check == "DUPLICATE_KEY"]
    assert f and f[0].status == "FAIL" and "rows 3, 5" in f[0].cells[0]


def test_composite_duplicate_key(tmp_path):
    rows = [["Store", "Article"], ["S1", "A"], ["S1", "B"], ["S1", "A"]]
    p = _book(tmp_path / "dup2.xlsx", rows)
    assert "DUPLICATE_KEY" in _checks(p, keys=[("Data", ["Store", "Article"])])


def test_case_variant_detected(tmp_path):
    p = _book(tmp_path / "case.xlsx", CLEAN + [["S4", "WEST", 1]])
    f = [x for x in xlsx_qc.scan(xlsx_qc.read_workbook(p)) if x.check == "CASE_VARIANT"]
    assert f and f[0].column == "Zone" and "'West'" in f[0].detail and "'WEST'" in f[0].detail


def test_spacing_variant_detected(tmp_path):
    p = _book(tmp_path / "space.xlsx", CLEAN + [["S4", "West ", 1], ["S5", "North  ", 1]])
    assert _checks(p).count("SPACE_VARIANT") == 2


def test_numeric_columns_not_flagged_as_variants(tmp_path):
    p = _book(tmp_path / "num.xlsx", [["Code", "NSV"], ["001", 1], ["1", 1]])
    assert "SPACE_VARIANT" not in _checks(p) and "CASE_VARIANT" not in _checks(p)


def test_allowed_values_catches_stray_label(tmp_path):
    rows = [["Channel"], ["MT"], ["GT"], ["METock"]]
    p = _book(tmp_path / "allowed.xlsx", rows)
    f = xlsx_qc.scan(xlsx_qc.read_workbook(p), allowed=[("Data", "Channel", {"MT", "GT"})])
    assert [x.detail.split("'")[1] for x in f if x.check == "NOT_ALLOWED"] == ["METock"]


def test_hidden_sheet_and_merged_cells_warned(tmp_path):
    def extra(wb):
        wb.active.merge_cells("A5:B5")
        wb.create_sheet("Secret").sheet_state = "hidden"
    p = _book(tmp_path / "struct.xlsx", CLEAN, extra=extra)
    c = _checks(p)
    assert "HIDDEN_SHEET" in c and "MERGED_CELLS" in c
    assert xlsx_qc.main([str(p)]) == 0          # WARN only -> still PASS


@pytest.mark.parametrize("name,content,code", [
    ("garbage.xlsx", b"this is not a workbook", "NOT_A_ZIP"),
    ("old.xls", b"\xd0\xcf\x11\xe0", "UNSUPPORTED_FORMAT"),
    ("binary.xlsb", b"PK\x03\x04", "UNSUPPORTED_FORMAT"),
])
def test_corrupt_or_unsupported_is_named_error(tmp_path, name, content, code):
    p = tmp_path / name
    p.write_bytes(content)
    with pytest.raises(xlsx_qc.WorkbookQCError) as e:
        xlsx_qc.read_workbook(p)
    assert e.value.code == code
    assert xlsx_qc.main([str(p)]) == 2


def test_zip_without_workbook_part(tmp_path):
    p = tmp_path / "empty.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(xlsx_qc.WorkbookQCError) as e:
        xlsx_qc.read_workbook(p)
    assert e.value.code == "MISSING_PART"


def test_entity_declaration_refused(tmp_path):
    p = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("xl/workbook.xml",
                   '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><workbook>&a;</workbook>')
    with pytest.raises(xlsx_qc.WorkbookQCError) as e:
        xlsx_qc.read_workbook(p)
    assert e.value.code == "UNSAFE_XML"


def test_missing_file_and_bad_argument(tmp_path):
    assert xlsx_qc.main([str(tmp_path / "nope.xlsx")]) == 2
    p = _book(tmp_path / "c.xlsx", CLEAN)
    assert xlsx_qc.main([str(p), "--key", "Data:No Such Column"]) == 2
    assert xlsx_qc.main([str(p), "--out", str(p)]) == 2        # never overwrite the workbook


def test_no_network_used(tmp_path, monkeypatch):
    def boom(*a, **k):
        raise AssertionError("network call attempted")
    monkeypatch.setattr(socket, "socket", boom)
    monkeypatch.setattr(socket, "create_connection", boom)
    p = _book(tmp_path / "n.xlsx", CLEAN + [["S4", "WEST", 1]])
    assert xlsx_qc.main([str(p), "--out", str(tmp_path / "r.csv")]) == 0


def test_source_workbook_not_mutated(tmp_path):
    p = _book(tmp_path / "m.xlsx", CLEAN + [["S2", "#REF!", 1]])
    before = (hashlib.sha256(p.read_bytes()).hexdigest(), p.stat().st_mtime_ns)
    xlsx_qc.main([str(p), "--key", "Data:Store Code", "--out", str(tmp_path / "r.json")])
    assert (hashlib.sha256(p.read_bytes()).hexdigest(), p.stat().st_mtime_ns) == before
    assert sorted(x.name for x in tmp_path.iterdir()) == ["m.xlsx", "r.json"]


def test_header_row_detected_below_title_banner(tmp_path):
    def extra(wb):
        wb.active.merge_cells("A1:C1")
    rows = [["MT Store Master | Jul-26"], ["Note: one row per store"],
            ["Store Code", "Zone", "NSV"], ["S1", "West", 1], ["S1", "WEST", 2]]
    p = _book(tmp_path / "banner.xlsx", rows, extra=extra)
    wb = xlsx_qc.read_workbook(p)
    assert wb["Data"].header_row == 3
    c = [f.check for f in xlsx_qc.scan(wb, keys=[("Data", ["Store Code"])])]
    assert "DUPLICATE_KEY" in c and "CASE_VARIANT" in c


def test_header_row_override(tmp_path):
    rows = [["a", "b"], ["Code", "Zone"], ["X", "West"], ["X", "East"]]
    p = _book(tmp_path / "ov.xlsx", rows)
    assert xlsx_qc.main([str(p), "--header-row", "Data:2", "--key", "Data:Code"]) == 1
    assert xlsx_qc.main([str(p), "--header-row", "Data:99"]) == 2
