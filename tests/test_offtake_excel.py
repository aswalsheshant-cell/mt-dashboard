"""Excel QoQ / H1 dashboard: structure and that the builder ties to the dashboard total (skips where the raw monthly CSVs are absent)."""
import importlib
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "PowerBI" / "RawDataFolders" / "Offtake_Monthly" / "offtake_store_article_Sep_26.csv"


@pytest.mark.skipif(not RAW.exists(), reason="raw monthly offtake CSVs are gitignored and not on this machine")
def test_workbook_builds_and_ties_to_dashboard(tmp_path):
    out = tmp_path / "x.xlsx"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "build_offtake_excel.py"), "--out", str(out)], check=True, capture_output=True)
    wb = openpyxl.load_workbook(out)
    for s in ("Dashboard", "Insights", "Control", "By_Chain", "By_Brand", "By_SubCategory", "By_PackSize", "Chain_x_Brand", "Monthly", "Pending_Chains", "Primary_vs_Offtake", "Offtake_Data"):
        assert s in wb.sheetnames
    for nm in ("d_Month", "d_NSV", "LatestMonth", "CY_H1_S", "LY_H1_E", "N_Q2", "d_Type", "Incl_Est", "Est_Crit"):
        assert nm in wb.defined_names
    ws = wb["By_Chain"]
    assert str(ws["B6"].value).startswith("=SUMIFS")   # numbers are formulas on the data sheet, not pasted values
    s = (ROOT / "dashboard" / "data.js").read_text()
    D = json.loads(s[s.index("{"): s.rindex("}") + 1])
    total = sum(r[6] for r in wb["Offtake_Data"].iter_rows(min_row=2, values_only=True) if r[0] and r[0].year == 2026 and r[0].month >= 4 and r[9] == "Actual")
    assert abs(total - D["offtake"]["total_fy27"]) < 0.05


# ---------- synthetic, formula-level tests (no confidential raw CSVs needed) ----------
import shutil
from datetime import date

sys.path.insert(0, str(ROOT / "scripts"))
be = importlib.import_module("build_offtake_excel")


def _synthetic(monkeypatch, months):
    cols = ["Month", "Chain", "Brand", "Category", "Sub-category", "Pack Size", "NSV", "Qty", "Source", "Type"]
    mon_no = {"Apr": 4, "May": 5, "Jun": 6, "Jul": 7, "Aug": 8, "Sep": 9}
    rows = []
    for m in months:
        for chain, v in (("ChainA", 100.0), ("ChainB", 50.0)):
            if chain == "ChainB" and m == months[-1]:
                continue                                    # ChainB has not reported the latest month -> estimate row
            rows.append([date(2026, mon_no[m], 1), chain, "Brand", "Cat", "Sub", "100 g/ml", v, 1.0, "POS", "Actual"])
    cy = pd.DataFrame(rows, columns=cols)
    ly = pd.DataFrame([[date(2025, n, 1), c, "n/a", "n/a", "n/a", "n/a", 40.0, 1.0, "POS", "Actual"] for n in mon_no.values() for c in ("ChainA", "ChainB")], columns=cols)
    monkeypatch.setattr(be, "available_months", lambda year="26": months)
    monkeypatch.setattr(be, "load_cy", lambda m, year="26": cy)
    monkeypatch.setattr(be, "load_ly", lambda: ly)
    return cy


def test_estimate_switch_defaults_to_actuals_only(monkeypatch, tmp_path):
    _synthetic(monkeypatch, ["Apr", "May", "Jun", "Jul"])
    out = tmp_path / "x.xlsx"
    monkeypatch.setattr(sys, "argv", ["build_offtake_excel.py", "--out", str(out)])
    be.main()
    wb = openpyxl.load_workbook(out)
    ref = wb.defined_names["Incl_Est"].attr_text
    sheet, cell = ref.split("!")
    assert wb[sheet.strip("'")][cell.replace("$", "")].value == 0        # default: actuals only
    assert "Contribution to" in "".join(str(c.value) for r in wb["By_Chain"].iter_rows(max_row=8) for c in r if c.value)
    assert "Share of total" not in "".join(str(c.value) for r in wb["By_Chain"].iter_rows(max_row=8) for c in r if c.value)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_recalculated_headline_excludes_estimate_and_counts_distinct_months(monkeypatch, tmp_path):
    _synthetic(monkeypatch, ["Apr", "Jun", "Jul", "Sep"])                  # May and Aug files missing
    out = tmp_path / "x.xlsx"
    monkeypatch.setattr(sys, "argv", ["build_offtake_excel.py", "--out", str(out)])
    be.main()
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "rc"), str(out)], check=True, capture_output=True, timeout=180)
    wb = openpyxl.load_workbook(tmp_path / "rc" / "x.xlsx", data_only=True)

    def named(n):
        sheet, cell = wb.defined_names[n].attr_text.split("!")
        return wb[sheet.strip("'")][cell.replace("$", "")].value

    assert named("N_H1") == 4                                              # four real months, not the six-month calendar span
    assert named("N_Q1") == 2                                              # Apr and Jun only (May missing)
    assert abs(named("CY_Total") - (100 + 100 + 100 + 100 + 50 + 50 + 50)) < 1e-6   # actuals only; the ChainB Sep estimate is left out
