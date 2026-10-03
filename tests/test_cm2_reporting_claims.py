"""Real temporary XLSX tests for signed workbook cuts and coverage."""
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import cm2_reporting as reporting


def workbook(tmp_path):
    w = openpyxl.Workbook()
    w.remove(w.active)
    sheets = {
        'Chain_Master': [['P&L chain', 'Format', 'TOT/MD %', 'Source of TOT/MD'], ['Apollo'], ['North RMTs'], ['WH-Smith']],
        'Chain_Alias': [['Source system', 'Chain name AS IT APPEARS in the source', 'P&L chain (locked name from your Provision sheet)', 'Note'], ['Primary', 'Apollo One', 'Apollo'], ['DN register', 'Apollo Two', 'Apollo'], ['DN register', 'WH Smith', 'WH-Smith']],
        'Data_Primary': [['Month', 'P&L chain (auto from Chain_Alias)', 'Brand', 'Primary NSV (Rs Lakh)', 'Chain name in source', 'Source'], ['Apr-26', '=NA()', 'Brand', 100, 'Apollo One', 'synthetic'], ['Apr-26', '=NA()', 'Brand', -20, 'Apollo Two', 'synthetic']],
        'Data_Provision': [['Month', 'P&L chain', 'Head', 'Amount (Rs Lakh)', 'Type', 'Your sheet cell', 'In Chain_Master?'], ['Apr-26', 'Apollo', 'Promotion', 10, 'Claim', 'O3', '=NA()'], ['Apr-26', 'Apollo', 'BA Expenses', 5, 'BA', 'V3', '=NA()'], ['Apr-26', 'North RMTs', 'Promotion', 7, 'Claim', 'O4', '=NA()']],
        'Data_ActualClaims': [['Month', 'P&L chain (auto from Chain_Alias)', 'Expense head', 'Actual claim - DN (Rs Lakh)', 'Chain name in source', 'Month status', 'Source'], ['Apr-26', '=NA()', 'Promotion', 12, 'Apollo One', 'Full', 'synthetic'], ['Apr-26', '=NA()', 'Credit', -2, 'Apollo Two', 'Full', 'synthetic'], ['Apr-26', '=NA()', 'Promotion', 3, 'Unknown', 'Full', 'synthetic'], ['Jul-26', '=NA()', 'Promotion', 2.2, 'WH Smith', 'PARTIAL MONTH', 'synthetic']],
    }
    for name, rows in sheets.items():
        s = w.create_sheet(name)
        for row in rows:
            s.append(row)
    path = tmp_path / ('synthetic-' + __import__('uuid').uuid4().hex + '.xlsx')
    w.save(path)
    return path


def test_signed_cuts_resolve_source_aliases_and_keep_ba_separate(tmp_path):
    claims, ba, dn = reporting.load_workbook_cuts(workbook(tmp_path))
    assert (len(claims), len(ba), len(dn)) == (2, 1, 4)
    assert [claims.amount_lakh.sum(), ba.amount_lakh.sum(), dn.amount_lakh.sum()] == pytest.approx([17, 5, 15.2])
    assert dn.chain.tolist() == ['Apollo', 'Apollo', 'UNMAPPED', 'WH-Smith']
    assert dn.source_status.tolist() == ['RECORDED_DN', 'RECORDED_DN', 'RECORDED_DN', 'PARTIAL_REGISTER']
    assert dn.mapping_status.tolist() == ['MAPPED', 'MAPPED', 'UNMAPPED', 'MAPPED']
    assert dn.source_row.tolist() == [2, 3, 4, 5]
    assert claims.source_status.tolist() == ['PROVISION', 'PROVISION']
    assert dn.month.tolist() == ['2026-04', '2026-04', '2026-04', '2026-07']
    assert dn.attrs['workbook_primary'].chain.tolist() == ['Apollo', 'Apollo']


def test_coverage_keeps_signed_sales_and_explicit_unallocated(tmp_path):
    claims, _, dn = reporting.load_workbook_cuts(workbook(tmp_path))
    primary = pd.DataFrame([['2026-04', 'Apollo', 100], ['2026-04', 'Apollo', -20], ['2026-04', 'Uncovered', -5], ['2026-07', 'WH-Smith', 0]], columns=['month', 'chain', 'nsv'])
    coverage = reporting.claim_coverage(primary, dn)
    assert coverage['covered_nsv'] == 80
    assert coverage['uncovered_nsv'] == -5
    assert coverage['source_total'] == pytest.approx(15.2)
    assert coverage['no_sales_claim_amount'] == pytest.approx(2.2)
    assert coverage['no_sales_keys'] == [('2026-07', 'WH-Smith')]
    assert coverage['unmapped_claim_amount'] == 3
    assert coverage['allocated_claim_amount'] == 10
    assert reporting.claim_coverage(primary, claims)['no_sales_claim_amount'] == 7
    assert coverage['unallocated'][0]['source_row'] == 4
    assert {r['status'] for r in coverage['unallocated']} == {'UNMAPPED', 'UNALLOCATED_NO_SALES'}


@pytest.mark.parametrize('sheet,column,value,match', [
    ('Data_ActualClaims', 6, 'Approved actual', 'status'),
    ('Data_Provision', 5, 'Actual', 'Type'),
    ('Data_Provision', 4, '=1+1', 'amount'),
    ('Data_ActualClaims', 1, 'bad month', 'month'),
])
def test_rejects_unsupported_or_formula_source_values(tmp_path, sheet, column, value, match):
    path = workbook(tmp_path)
    w = openpyxl.load_workbook(path)
    w[sheet].cell(2, column, value)
    w.save(path)
    with pytest.raises(ValueError, match=match):
        reporting.load_workbook_cuts(path)


def test_rejects_duplicate_provision_source_event(tmp_path):
    path = workbook(tmp_path)
    w = openpyxl.load_workbook(path)
    w['Data_Provision'].append([c.value for c in w['Data_Provision'][2]])
    w.save(path)
    with pytest.raises(ValueError, match='Duplicate.*Data_Provision.*5'):
        reporting.load_workbook_cuts(path)


def test_rejects_duplicate_dn_source_event(tmp_path):
    path = workbook(tmp_path)
    w = openpyxl.load_workbook(path)
    w['Data_ActualClaims'].append([c.value for c in w['Data_ActualClaims'][2]])
    w.save(path)
    with pytest.raises(ValueError, match='Duplicate.*Data_ActualClaims.*6'):
        reporting.load_workbook_cuts(path)


@pytest.mark.parametrize('fault', ['sheet', 'column', 'alias'])
def test_rejects_missing_contract_and_conflicting_aliases(tmp_path, fault):
    path = workbook(tmp_path)
    w = openpyxl.load_workbook(path)
    if fault == 'sheet':
        del w['Data_Primary']
    elif fault == 'column':
        w['Data_Primary'].cell(1, 5, 'Wrong')
    else:
        w['Chain_Alias'].append(['DN register', 'Apollo One', 'WH-Smith'])
    w.save(path)
    with pytest.raises(ValueError):
        reporting.load_workbook_cuts(path)
@pytest.fixture
def tmp_path():
    # Use files in an existing ignored directory because this Windows
    # sandbox restricts new mode-0700 directories used by pytest.
    folder = Path(__file__).resolve().parents[1] / '.superpowers'
    before = set(folder.glob('synthetic-*.xlsx'))
    yield folder
    for path in set(folder.glob('synthetic-*.xlsx')) - before:
        path.unlink()

def test_provision_no_sales_control_keeps_claim_and_ba_split(tmp_path):
    path = workbook(tmp_path)
    w = openpyxl.load_workbook(path)
    w['Data_Provision'].cell(4, 4, 69.4529)
    w['Data_Provision'].append(['Apr-26', 'North RMTs', 'BA Expenses', 68.5, 'BA', 'V4', '=NA()'])
    w.save(path)
    claims, ba, _ = reporting.load_workbook_cuts(path)
    primary = claims.attrs['workbook_primary']
    claim_unallocated = reporting.claim_coverage(primary, claims)['no_sales_claim_amount']
    ba_unallocated = reporting.claim_coverage(primary, ba)['no_sales_claim_amount']
    assert claim_unallocated == pytest.approx(69.4529)
    assert ba_unallocated == pytest.approx(68.5)
    assert claim_unallocated + ba_unallocated == pytest.approx(137.9529)


def test_net_zero_sales_stay_unallocated_and_negative_sales_are_covered(tmp_path):
    _, _, dn = reporting.load_workbook_cuts(workbook(tmp_path))
    primary = pd.DataFrame([['2026-04', 'Apollo', -20], ['2026-07', 'WH-Smith', 4], ['2026-07', 'WH-Smith', -4]], columns=['month', 'chain', 'nsv'])
    coverage = reporting.claim_coverage(primary, dn)
    assert coverage['covered_nsv'] == -20
    assert coverage['uncovered_nsv'] == 0
    assert coverage['no_sales_claim_amount'] == pytest.approx(2.2)
    assert coverage['no_sales_keys'] == [('2026-07', 'WH-Smith')]
