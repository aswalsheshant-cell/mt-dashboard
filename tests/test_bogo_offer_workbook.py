"""BOGO offer-wise offtake workbook builder: rules, helpers and reconciliation to the source files."""
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import build_bogo_offer_workbook as W  # noqa: E402

OFFTAKE = list((ROOT / "PowerBI" / "RawDataFolders" / "Offtake_Monthly").glob("offtake_store_article_*_*.csv"))


def test_helpers():
    assert W.mlabel("Apr_26") == "Apr-26" and W.mlabel("2026-09") == "Sep-26"
    assert W.ean("'8904417306255") == "8904417306255" and W.ean(8.904417e12) == "8904417000000" and W.ean(float("nan")) == ""
    assert W._offer_text("BOGO") == ("BOGO", "BOGO")
    assert W._offer_text("0.25") == ("25% off", "Discount %")
    assert W._offer_text("Not Specified")[1] == "Not specified"
    assert W._offer_text(float("nan")) == ("", "")


@pytest.mark.skipif(not OFFTAKE, reason="offtake source files not present")
def test_workbook_reconciles_and_keeps_eligibility_apart_from_offers(tmp_path):
    out = tmp_path / "w.xlsx"
    info = W.build(str(out), 0.55)
    wb = openpyxl.load_workbook(out, data_only=True)
    assert {"Raw_Data", "Checks", "Review_List", "Sum_Chain", "Sum_Pack", "Offer_Compare", "MoM_Trends"} <= set(wb.sheetnames)
    chk = wb["Checks"]
    status = [r[5] for r in chk.iter_rows(min_row=5, values_only=True) if r[0] and str(r[0]).startswith(("Offtake value, excl", "Offtake units, excl", "Sum_", "Chain_Month"))]
    assert status and all(s == "OK" for s in status), status
    rows = list(wb["Raw_Data"].iter_rows(values_only=True))
    h = rows[0]
    el, mg, offer = h.index("BOGO eligibility"), h.index("Article margin %"), h.index("Promotional offer communicated")
    for r in rows[1:]:
        if r[el] == "BOGO eligible":
            assert r[mg] is not None and r[mg] >= 0.55
        if r[el] == "Non-BOGO":
            assert r[mg] is not None and 0 < r[mg] < 0.55
        if r[offer] in (None, ""):
            assert "No offer communication on file" in (r[h.index("Review flag")] or "")   # blank, flagged, never guessed
    assert info["offers_matched"] >= 0
