"""The monthly deck must never present sample or typed-in numbers as real.

Found 2026-10-14: monthly_mt_deck.yml ran on a schedule (run #1, 2026-10-01) with
no data source, so it fell back to data/sample_seeds plus the typed-in DEFAULT_CONFIG
in scripts/build_mt_monthly_ppt.py, and produced "MT-Leadership-Deck-september-2026"
with no sign it was sample data. Same workflow: the summary always printed
"Passed (32/32 tests)", a missing data path was skipped silently (build fell back to
defaults), inputs were pasted into shell commands, and the Slides deploy passed
--credentials and --share, which deploy_to_google_slides.py does not accept (it only
does interactive OAuth), so it could never run in CI.

These tests pin the fix. The deck is not fed from dashboard/data.js, so every deck
the CLI builds is stamped SAMPLE or DRAFT.
"""
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent
WF = ROOT / ".github" / "workflows" / "monthly_mt_deck.yml"
SCRIPTS = ROOT / "scripts"
BUILDER = SCRIPTS / "build_mt_monthly_ppt.py"
SEEDS = ROOT / "data" / "sample_seeds"


def _wf():
    return yaml.safe_load(WF.read_text(encoding="utf-8"))


def _code_lines():
    return [ln for ln in WF.read_text(encoding="utf-8").splitlines() if not ln.lstrip().startswith("#")]


def test_workflow_has_no_schedule_trigger():
    wf = _wf()
    triggers = wf.get("on", wf.get(True))
    assert "schedule" not in triggers, (
        "A scheduled run has no data source and builds from sample data. "
        "Re-add a schedule only after the builder reads verified data (dashboard/data.js)."
    )


def test_workflow_never_defaults_to_sample_seeds_silently():
    code = "\n".join(_code_lines())
    assert "|| './data/sample_seeds'" not in code
    assert "github.event_name == 'schedule'" not in code


def test_workflow_does_not_fake_a_test_result():
    assert not re.search(r"Passed \(\d+/\d+", "\n".join(_code_lines()))


def test_workflow_inputs_are_not_pasted_into_shell():
    for job in _wf()["jobs"].values():
        for step in job["steps"]:
            run = step.get("run", "")
            assert "github.event.inputs" not in run, step.get("name")
            assert "eval " not in run, step.get("name")


def test_workflow_passes_only_flags_the_scripts_support():
    code = "\n".join(_code_lines())
    for script in ("build_mt_monthly_ppt.py", "deploy_to_google_slides.py"):
        for m in re.finditer(r"scripts/" + re.escape(script) + r"([^\n]*(?:\\\n[^\n]*)*)", code):
            flags = set(re.findall(r"(--[a-z][a-z-]*)", m.group(1)))
            src = (SCRIPTS / script).read_text(encoding="utf-8")
            for flag in flags:
                assert f'"{flag}"' in src, f"{script} has no {flag}"


def test_workflow_missing_data_source_fails():
    code = "\n".join(_code_lines())
    assert "Data source not found" in code and "exit 1" in code


@pytest.fixture(scope="module")
def builder():
    pytest.importorskip("pptx")
    sys.path.insert(0, str(SCRIPTS))
    try:
        import build_mt_monthly_ppt as mod
    finally:
        sys.path.remove(str(SCRIPTS))
    return mod


def test_classify_data_status(builder):
    assert builder.classify_data_status(None) == "SAMPLE"
    assert builder.classify_data_status("") == "SAMPLE"
    assert builder.classify_data_status("./data/sample_seeds") == "SAMPLE"
    assert builder.classify_data_status("data/sample_seeds/") == "SAMPLE"
    assert builder.classify_data_status("some/other/dir") == "DRAFT"
    assert builder.classify_data_status("metrics.json") == "DRAFT"


def _run(args, cwd):
    return subprocess.run([sys.executable, str(BUILDER), *args], cwd=cwd,
                          capture_output=True, text=True)


def _slide_texts(pptx_path):
    from pptx import Presentation
    prs = Presentation(str(pptx_path))
    return [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]


@pytest.mark.parametrize("source_name,status,banner", [
    ("sample_seeds", "SAMPLE", "SAMPLE DATA - NOT FOR DISTRIBUTION"),
    ("real_seeds", "DRAFT", "DRAFT - NOT VERIFIED AGAINST dashboard/data.js"),
])
def test_every_slide_and_json_are_stamped(builder, tmp_path, source_name, status, banner):
    src = tmp_path / source_name
    shutil.copytree(SEEDS, src)
    out = tmp_path / "deck"
    r = _run(["--month", "september", "--year", "2026", "--format", "both",
              "--output", str(out), "--data-source", str(src)], tmp_path)
    assert r.returncode == 0, r.stderr
    texts = _slide_texts(f"{out}.pptx")
    assert texts and all(banner in t for t in texts)
    payload = json.loads(Path(f"{out}_gslides_batch.json").read_text(encoding="utf-8"))
    assert f"[{status}]" in json.dumps(payload)


def test_print_data_status_exits_without_building(builder, tmp_path):
    r = _run(["--print-data-status", "--data-source", str(SEEDS)], tmp_path)
    assert r.returncode == 0 and r.stdout.strip() == "SAMPLE"
    assert not list(tmp_path.iterdir())


def test_missing_data_source_does_not_fall_back(builder, tmp_path):
    r = _run(["--format", "pptx", "--output", str(tmp_path / "x"),
              "--data-source", str(tmp_path / "nope")], tmp_path)
    assert r.returncode != 0
    assert not (tmp_path / "x.pptx").exists()


def test_other_callers_are_unchanged(builder, tmp_path):
    cfg = dict(builder.DEFAULT_CONFIG)
    cfg.pop("data_banner", None)
    out = tmp_path / "plain.pptx"
    builder.build_mt_ppt(cfg, str(out))
    assert not any("NOT FOR DISTRIBUTION" in t or "NOT VERIFIED" in t for t in _slide_texts(out))
